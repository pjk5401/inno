import base64
import io
import json
import math
import shutil
import zipfile
import uuid
from datetime import timedelta
from functools import wraps
from pathlib import Path
import openpyxl
import qrcode
from openpyxl.drawing.image import Image as ExcelImage
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout, get_user_model
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import connection, transaction
from django.db.models import Q, Count
from django.http import JsonResponse, HttpResponse, FileResponse, Http404
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST, require_GET
from .models import Policy, Course, Enrollment, Attempt, Attendance, Audit, BrowserIdentity
from .forms import CourseForm, LocationForm, PolicyForm, EnrollmentForm, CorrectionForm
from . import services as svc

def staff_required(view=None, *, admin=False):
    def decorator(fn):
        @wraps(fn)
        def inner(request,*args,**kwargs):
            if not request.user.is_authenticated:
                return redirect('login')
            if not request.user.is_active or not request.user.is_staff or (admin and not request.user.is_superuser):
                return render(request,'error.html',{'error':'이 기능을 사용할 권한이 없습니다.'},status=403)
            return fn(request,*args,**kwargs)
        return inner
    return decorator(view) if view else decorator

def browser_id(request):
    value = request.get_signed_cookie(settings.BROWSER_COOKIE_NAME, default='', salt='browser-identity', max_age=settings.BROWSER_COOKIE_AGE)
    try:
        return uuid.UUID(value)
    except (ValueError, TypeError):
        return None

def get_attempt(request):
    value = request.session.get('attempt')
    if not value:
        return None
    try:
        identity = browser_id(request)
        if identity is None:
            return None
        return Attempt.objects.select_related('completed__enrollment__course','browser').filter(pk=value,browser_id=identity).first()
    except (ValidationError,ValueError):
        return None

def home(request):
    return render(request,'home.html')

def csrf_failure(request, reason=''):
    return render(request,'error.html',{'error':'접속 정보가 만료되었거나 요청 출처를 확인하지 못했습니다. 화면을 새로고침하거나 최신 QR로 다시 접속해주세요.'},status=403)

def enter(request):
    try:
        attempt = svc.exchange_qr(request.GET.get('q',''), browser_id(request))
        request.session['attempt'] = str(attempt.id)
        response = redirect('checkin')
        response.set_signed_cookie(settings.BROWSER_COOKIE_NAME, str(attempt.browser_id), salt='browser-identity',
            max_age=settings.BROWSER_COOKIE_AGE, path=settings.SESSION_COOKIE_PATH,
            secure=not settings.DEBUG, httponly=True, samesite='Lax')
        return response
    except svc.Rejected as e:
        return render(request,'error.html',{'error':str(e)},status=400)

def checkin(request):
    p = Policy.current()
    attempt = get_attempt(request)
    try:
        svc.valid_attempt(attempt,p)
    except svc.Rejected as e:
        return render(request,'error.html',{'error':str(e)},status=400)
    if attempt.completed_id:
        return redirect('complete')
    at = svc.now(p)
    day = timezone.localtime(at).date()
    courses = list(Course.objects.filter(Q(starts_at__date=day) |
        Q(starts_at__lte=at+timedelta(minutes=p.before_minutes),ends_at__gt=at)))
    for c in courses:
        try:
            c.available = True
            c.state = {'present':'출석 접수','late':'지각 접수'}[svc.course_status(c,p,at)]
        except svc.Rejected as e:
            c.available,c.state = False, str(e)
    return render(request,'checkin.html',{'courses':courses,'policy':p,
        'remaining':max(0,int((attempt.expires_at-at).total_seconds())),
        'generation':p.generation,'identity':attempt.browser,'mock':settings.SANDBOX and p.location_mode!='real'})

@require_POST
def submit(request):
    attempt = get_attempt(request)
    try:
        svc.throttle('submit:'+str(attempt.id if attempt else request.session.session_key),30)
        payload = json.loads(request.body)
        if not isinstance(payload,dict):
            raise ValueError()
        record = svc.submit(attempt.id if attempt else None,payload)
        return JsonResponse({'ok':True,'redirect':reverse('complete')})
    except (json.JSONDecodeError,ValueError) as e:
        message = str(e) if isinstance(e,svc.Rejected) else '입력 형식이 올바르지 않습니다.'
        return JsonResponse({'ok':False,'error':message},status=400)

def attempt_status(request):
    p = Policy.current()
    attempt = get_attempt(request)
    valid = bool(attempt and attempt.generation == p.generation and attempt.browser_id and attempt.browser_revision == attempt.browser.revision)
    remaining = max(0,int((attempt.expires_at-svc.now(p)).total_seconds())) if valid else 0
    return JsonResponse({'valid':valid,'remaining':remaining,'generation':p.generation})

def complete(request):
    attempt = get_attempt(request)
    if not attempt or attempt.generation != Policy.current().generation or not attempt.completed_id or not attempt.browser_id or attempt.browser_revision != attempt.browser.revision:
        return render(request,'error.html',{'error':'출석 완료 기록을 확인할 수 없습니다. 조교에게 문의해주세요.'},status=404)
    return render(request,'complete.html',{'record':attempt.completed,'identity':attempt.browser})

def staff_login(request):
    if request.method == 'POST':
        username = request.POST.get('username','').strip()[:150]
        try:
            svc.throttle('login:'+request.META.get('REMOTE_ADDR','')+':'+username,8,300)
            user = authenticate(request,username=username,password=request.POST.get('password',''))
            if user and user.is_active and user.is_staff:
                login(request,user)
                return redirect('dashboard')
            messages.error(request,'아이디 또는 비밀번호를 확인해주세요.')
        except svc.Rejected as e:
            messages.error(request,str(e))
    return render(request,'login.html')

@staff_required
def browser_identity(request):
    code = (request.POST.get('code','') if request.method == 'POST' else request.GET.get('code','')).strip().upper()
    identity = BrowserIdentity.objects.filter(support_code=code).first() if code else None
    if request.method == 'POST':
        try:
            sid = request.POST.get('student_id','').strip()
            name = request.POST.get('name','').strip()
            reason = request.POST.get('reason','').strip()
            if not sid or not name or not reason or len(reason)>1000:
                raise svc.Rejected('학번·이름과 변경 사유(1,000자 이내)를 입력해주세요.')
            with transaction.atomic():
                Policy.objects.select_for_update().get(pk=Policy.current().pk)
                identity = BrowserIdentity.objects.select_for_update().filter(support_code=code).first()
                if not identity:
                    raise svc.Rejected('브라우저 확인 코드를 찾을 수 없습니다.')
                if str(identity.revision) != request.POST.get('revision'):
                    raise svc.Rejected('다른 담당자가 변경했습니다. 다시 조회해주세요.')
                if not Enrollment.objects.filter(student_id=sid,name=name).exists():
                    raise svc.Rejected('명단에 등록된 학번·이름을 입력해주세요.')
                before = {'student_id':identity.student_id,'name':identity.name}
                identity.student_id, identity.name = sid, name
                identity.revision += 1
                identity.save()
                svc.audit(request.user,'브라우저 학번 변경',identity.support_code,before=before,
                          after={'student_id':sid,'name':name},reason=reason)
            messages.success(request,'학번을 변경했습니다. 학생은 새 QR로 다시 접속해주세요.')
            return redirect(reverse('browser_identity')+'?code='+code)
        except svc.Rejected as e:
            messages.error(request,str(e))
    elif code and not identity:
        messages.error(request,'브라우저 확인 코드를 찾을 수 없습니다.')
    return render(request,'browser_identity.html',{'identity':identity,'code':code})

@require_POST
def staff_logout(request):
    logout(request)
    return redirect('login')

@staff_required
def dashboard(request):
    p = Policy.current()
    date = request.GET.get('date') or timezone.localtime(svc.now(p)).date().isoformat()
    try:
        from datetime import date as date_type
        date_type.fromisoformat(date)
    except ValueError:
        date = timezone.localtime(svc.now(p)).date().isoformat()
    courses = Course.objects.filter(starts_at__date=date).annotate(
        total=Count('enrollments'),present=Count('enrollments',filter=Q(enrollments__attendance__status='present')),
        late=Count('enrollments',filter=Q(enrollments__attendance__status='late')))
    for c in courses:
        c.missing = c.total-c.present-c.late
        c.ended = svc.now(p)>=c.ends_at
    return render(request,'dashboard.html',{'courses':courses,'date':date,'policy':p,'clock':svc.now(p)})

@staff_required
def courses(request):
    return render(request,'courses.html',{'courses':Course.objects.all()})

@staff_required
def course_edit(request,pk=None):
    instance = get_object_or_404(Course,pk=pk) if pk else None
    before = {'title':instance.title,'starts_at':instance.starts_at.isoformat(),'ends_at':instance.ends_at.isoformat(),'active':instance.active} if instance else {}
    form = CourseForm(request.POST or None,instance=instance)
    if request.method=='POST' and form.is_valid():
        c=form.save()
        svc.audit(request.user,'수업 저장',c.pk,before=before,after={'title':c.title,'starts_at':c.starts_at.isoformat(),'ends_at':c.ends_at.isoformat(),'active':c.active})
        messages.success(request,'수업을 저장했습니다. 기존 출석 판정은 변경하지 않습니다.')
        return redirect('course',pk=c.pk)
    return render(request,'form.html',{'form':form,'title':'수업 수정' if pk else '수업 등록','note':'기록이 있는 수업은 삭제하지 않고 접수 사용을 해제할 수 있습니다.'})

@staff_required
def course_detail(request,pk):
    c=get_object_or_404(Course,pk=pk)
    q=request.GET.get('q','').strip()
    state=request.GET.get('state','')
    rows=list(c.enrollments.select_related('attendance','course').filter(Q(name__icontains=q)|Q(student_id__icontains=q)|Q(department__icontains=q)))
    p=Policy.current()
    for row in rows:
        row.state=svc.display_status(row,p)
    if state:
        rows=[r for r in rows if r.state==state]
    page=Paginator(rows,100).get_page(request.GET.get('page'))
    return render(request,'course.html',{'course':c,'rows':page,'q':q,'state':state})

@staff_required
def enrollment_edit(request,course_id,pk=None):
    c=get_object_or_404(Course,pk=course_id)
    instance=get_object_or_404(Enrollment,pk=pk,course=c) if pk else Enrollment(course=c)
    before={'student_id':instance.student_id,'name':instance.name,'department':instance.department}
    form=EnrollmentForm(request.POST or None,instance=instance)
    if request.method=='POST' and form.is_valid():
        if Enrollment.objects.filter(course=c,student_id=form.cleaned_data['student_id']).exclude(pk=instance.pk).exists():
            form.add_error('student_id','해당 수업에 이미 등록된 학번입니다.')
        else:
            with transaction.atomic():
                row=form.save()
                svc.audit(request.user,'명단 수정',row.pk,before=before,after=form.cleaned_data)
            return redirect('course',pk=c.pk)
    return render(request,'form.html',{'form':form,'title':'학생 명단 수정' if pk else '학생 추가','note':'학생 계정은 생성하지 않습니다. 기존 서명 당시 이름은 출석 상세에 별도로 보존됩니다.'})

def read_xlsx(upload):
    if not upload or not upload.name.lower().endswith('.xlsx') or upload.size>5*1024*1024:
        raise svc.Rejected('5MB 이하의 .xlsx 파일을 선택해주세요.')
    raw=upload.read()
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(x.file_size for x in archive.infolist())>30*1024*1024:
                raise svc.Rejected('압축을 해제한 엑셀 크기가 너무 큽니다.')
        book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        sheet=book.active
        rows=[]
        for i,row in enumerate(sheet.iter_rows()):
            if i>10000 or len(row)>50:
                raise svc.Rejected('명단은 최대 10,000행, 50열까지 등록할 수 있습니다.')
            values=[]
            for cell in row:
                if cell.data_type=='f':
                    raise svc.Rejected('수식이 있는 파일입니다. 값을 붙여넣은 파일로 등록해주세요.')
                value=cell.value
                if isinstance(value,(int,float)) and not isinstance(value,bool):
                    if isinstance(value,float) and not value.is_integer():
                        value=str(value)
                    elif cell.number_format and set(cell.number_format)=={'0'}:
                        value=str(int(value)).zfill(len(cell.number_format))
                    else:
                        value=str(int(value))
                values.append('' if value is None else str(value).strip())
            if any(values):
                rows.append(values)
        book.close()
        if len(rows)<2:
            raise svc.Rejected('첫 행에 제목, 다음 행부터 학생 정보를 넣어주세요.')
        return rows
    except (zipfile.BadZipFile,KeyError,ValueError,OSError) as e:
        raise svc.Rejected(str(e) if isinstance(e,svc.Rejected) else '읽을 수 없는 엑셀 파일입니다.')

@staff_required
def roster_import(request,pk):
    c=get_object_or_404(Course,pk=pk)
    key=f'import_{pk}'
    pending=request.session.get(key)
    if request.method=='POST':
        try:
            action=request.POST.get('action')
            if action=='upload':
                rows=read_xlsx(request.FILES.get('file'))
                pending={'headers':rows[0],'rows':rows[1:]}
                request.session[key]=pending
            elif action=='map' and pending:
                indices={k:int(request.POST[k]) for k in ['department','student_id','name']}
                if len(set(indices.values()))<3:
                    raise svc.Rejected('학과·학번·이름은 서로 다른 열을 선택해주세요.')
                output=[]; seen=set(); errors=[]
                for i,row in enumerate(pending['rows'],2):
                    item={k:row[n] if 0<=n<len(row) else '' for k,n in indices.items()}
                    if not item['student_id'] or not item['name']:
                        errors.append(f'{i}행: 학번·이름이 필요합니다.')
                    if len(item['student_id'])>40 or len(item['name'])>80 or len(item['department'])>100:
                        errors.append(f'{i}행: 입력 길이가 너무 깁니다.')
                    if item['student_id'] in seen:
                        errors.append(f'{i}행: 중복 학번 {item["student_id"]}')
                    seen.add(item['student_id']); output.append(item)
                pending['parsed']=output
                pending['errors']=errors[:50]
                existing={e.student_id:e for e in c.enrollments.all()}
                pending['add_count']=sum(x['student_id'] not in existing for x in output)
                pending['update_count']=sum(x['student_id'] in existing and (x['name']!=existing[x['student_id']].name or x['department']!=existing[x['student_id']].department) for x in output)
                request.session[key]=pending
            elif action=='confirm' and pending and pending.get('parsed') and not pending.get('errors'):
                with transaction.atomic():
                    changes=[]
                    for row in pending['parsed']:
                        e,created=Enrollment.objects.get_or_create(course=c,student_id=row['student_id'],defaults={'name':row['name'],'department':row['department']})
                        if not created and (e.name!=row['name'] or e.department!=row['department']):
                            changes.append({'id':e.pk,'before':{'name':e.name,'department':e.department},'after':row})
                            e.name,e.department=row['name'],row['department'];e.save()
                    svc.audit(request.user,'명단 업로드',c.pk,count=len(pending['parsed']),changes=changes)
                del request.session[key]
                messages.success(request,'명단을 등록했습니다. 파일에 없는 기존 학생과 출석 기록은 유지됩니다.')
                return redirect('course',pk=c.pk)
            else:
                raise svc.Rejected('파일을 다시 선택해주세요.')
        except (svc.Rejected,ValueError,KeyError) as e:
            messages.error(request,str(e))
    return render(request,'import.html',{'course':c,'pending':pending})

def revision(record):
    if not record:
        return 'new'
    import hashlib
    return hashlib.sha256(json.dumps({'status':record.status,'at':str(record.attended_at),'reason':record.reason,'method':record.method},sort_keys=True).encode()).hexdigest()

@staff_required
def record_detail(request,pk):
    row=get_object_or_404(Enrollment.objects.select_related('course'),pk=pk)
    record=Attendance.objects.filter(enrollment=row).first()
    initial={'status':record.status if record else 'present','attended_at':timezone.localtime(record.attended_at if record and record.attended_at else svc.now()),'revision':revision(record)}
    form=CorrectionForm(request.POST or None,initial=initial)
    if request.method=='POST' and form.is_valid():
        with transaction.atomic():
            Enrollment.objects.select_for_update().get(pk=pk)
            record=Attendance.objects.filter(enrollment=row).first()
            if form.cleaned_data['revision']!=revision(record):
                form.add_error(None,'다른 작업에서 기록이 변경되었습니다. 새로고침 후 다시 확인해주세요.')
            else:
                before={'status':record.status,'attended_at':str(record.attended_at),'reason':record.reason} if record else {}
                if not record:
                    record=Attendance(enrollment=row)
                record.status=form.cleaned_data['status'];record.attended_at=form.cleaned_data['attended_at']
                record.method='manual';record.reason=form.cleaned_data['reason'];record.save()
                svc.audit(request.user,'출석 정정',row.pk,before=before,after={'status':record.status,'attended_at':str(record.attended_at),'reason':record.reason})
                messages.success(request,'사유와 변경 이력을 저장했습니다.')
                return redirect('record',pk=pk)
    history=Audit.objects.filter(target=str(pk),action__in=['출석 정정','명단 수정'])[:30]
    return render(request,'record.html',{'row':row,'record':record,'form':form,'history':history})

@staff_required
def signature(request,pk):
    record=get_object_or_404(Attendance,pk=pk)
    if not record.signature:
        raise Http404()
    return FileResponse(record.signature.open('rb'),content_type='image/png')

def safe_cell(value):
    if isinstance(value,str) and value.startswith(('=','+','-','@','\t','\r')):
        return "'"+value
    return value

@staff_required
def export(request):
    courses=Course.objects.all()
    if request.GET.get('course'):
        try:
            course_id=int(request.GET['course'])
        except ValueError:
            return HttpResponse('수업 번호를 확인해주세요.',status=400)
        courses=courses.filter(pk=course_id)
    if request.GET.get('date'):
        try:
            from datetime import date
            date.fromisoformat(request.GET['date'])
        except ValueError:
            return HttpResponse('날짜 형식을 확인해주세요.',status=400)
        courses=courses.filter(starts_at__date=request.GET['date'])
    rows=Enrollment.objects.filter(course__in=courses).select_related('course','attendance').order_by('course__starts_at','course_id','student_id')
    include=request.GET.get('signatures')=='1'
    book=openpyxl.Workbook();sheet=book.active;sheet.title='출석부'
    sheet.append(['일자','수업명','분반','학과','학번','이름','출석상태','출석시간','처리방식','환경','서명시 이름']+(['서명'] if include else []))
    p=Policy.current()
    for num,row in enumerate(rows,2):
        a=getattr(row,'attendance',None)
        values=[timezone.localtime(row.course.starts_at).date().isoformat(),row.course.title,row.course.section,row.department,row.student_id,row.name,svc.display_status(row,p),timezone.localtime(a.attended_at).strftime('%Y-%m-%d %H:%M:%S') if a and a.attended_at else '',a.get_method_display() if a else '',('테스트 · '+(a.location_source if a else '')) if settings.SANDBOX else '운영',a.signed_name if a else '']
        sheet.append([safe_cell(x) for x in values])
        sheet.cell(num,5).number_format='@'
        if include and a and a.signature:
            with a.signature.open('rb') as f:
                image=ExcelImage(io.BytesIO(f.read()))
            ratio=min(190/image.width,65/image.height)
            image.width*=ratio;image.height*=ratio
            sheet.add_image(image,f'L{num}');sheet.row_dimensions[num].height=55
    from openpyxl.styles import Font, PatternFill
    for cell in sheet[1]:
        cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='176A53')
    for col,width in {'A':14,'B':26,'C':12,'D':18,'E':18,'F':16,'G':12,'H':23,'I':14,'J':22,'K':18,'L':29}.items():
        sheet.column_dimensions[col].width=width
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    out=io.BytesIO();book.save(out);out.seek(0)
    svc.audit(request.user,'출석부 다운로드',request.GET.get('course','all'),signatures=include,date=request.GET.get('date',''))
    return FileResponse(out,as_attachment=True,filename=('test-' if settings.SANDBOX else '')+('attendance-signatures.xlsx' if include else 'attendance.xlsx'))

@staff_required
def location(request):
    p=Policy.current();before=svc.snapshot(p)
    form=LocationForm(request.POST or None,instance=p)
    if request.method=='POST' and form.is_valid():
        p=form.save();svc.audit(request.user,'인정 구역 변경','policy',before=before,after=svc.snapshot(p))
        messages.success(request,'인정 구역을 저장했습니다. 이후 제출부터 적용합니다.');return redirect('location')
    return render(request,'location.html',{'form':form,'policy':p})

@staff_required
def qr_display(request):
    return render(request,'qr.html',{'expired':settings.SANDBOX and request.GET.get('expired')=='1'})

@staff_required
def qr_data(request):
    p=Policy.current();at=svc.now(p)
    expired=settings.SANDBOX and request.GET.get('expired')=='1'
    token=svc.qr_token(p,expired)
    from urllib.parse import urlencode
    url=settings.PUBLIC_ORIGIN+reverse('enter')+'?'+urlencode({'q':token})
    buf=io.BytesIO();qrcode.make(url,box_size=8,border=3).save(buf,format='PNG')
    return JsonResponse({'image':'data:image/png;base64,'+base64.b64encode(buf.getvalue()).decode(),
        'url':url,'remaining':p.qr_seconds-int(at.timestamp())%p.qr_seconds,'clock':timezone.localtime(at).strftime('%Y-%m-%d %H:%M:%S'),
        'mode':p.get_location_mode_display() if settings.SANDBOX else '', 'expired':expired})

@staff_required
def audit_log(request):
    rows=Audit.objects.select_related('actor')
    q=request.GET.get('q','')
    if q:
        rows=rows.filter(Q(action__icontains=q)|Q(target__icontains=q)|Q(actor__username__icontains=q))
    return render(request,'audit.html',{'rows':Paginator(rows,50).get_page(request.GET.get('page')),'q':q})

@staff_required(admin=True)
def system(request):
    p=Policy.current();before=svc.snapshot(p)
    form=PolicyForm(request.POST or None,instance=p)
    if request.method=='POST' and form.is_valid():
        p=form.save();svc.audit(request.user,'출석 정책 변경','policy',before=before,after=svc.snapshot(p))
        messages.success(request,'정책을 저장했습니다. 기존 출석과 이미 발급된 작성 세션은 유지됩니다.');return redirect('system')
    Path(settings.MEDIA_ROOT).mkdir(parents=True,exist_ok=True)
    disk=shutil.disk_usage(settings.MEDIA_ROOT)
    backup=Path('/backups')
    backups=[]
    if backup.exists():
        backups=[{'name':x.name,'size':round(x.stat().st_size/1024/1024,1)} for x in sorted(backup.glob('*.tar.gz'),reverse=True)[:10]]
    return render(request,'system.html',{'form':form,'disk_free':round(disk.free/1024**3,1),'disk_total':round(disk.total/1024**3,1),'backups':backups})

@staff_required(admin=True)
def accounts(request):
    User=get_user_model()
    if request.method=='POST':
        try:
            action=request.POST.get('action')
            with transaction.atomic():
                if action=='create':
                    name=request.POST.get('username','').strip()
                    import re
                    if not re.fullmatch(r'[a-zA-Z0-9_.@+-]{1,150}',name):
                        raise svc.Rejected('아이디는 영문·숫자 및 . @ + - _를 사용해주세요.')
                    if User.objects.filter(username=name).exists():
                        raise svc.Rejected('이미 존재하는 계정입니다.')
                    user=User(username=name,is_staff=True,is_superuser=request.POST.get('role')=='admin')
                    validate_password(request.POST.get('password',''),user)
                    user.set_password(request.POST['password']);user.save()
                    svc.audit(request.user,'계정 생성',user.username,role='admin' if user.is_superuser else 'assistant')
                else:
                    user=get_object_or_404(User.objects.select_for_update(),pk=request.POST.get('user'))
                    if action=='password':
                        validate_password(request.POST.get('password',''),user)
                        user.set_password(request.POST['password']);user.save()
                        svc.audit(request.user,'비밀번호 초기화',user.username)
                    elif action=='access':
                        if user.pk==request.user.pk:
                            raise svc.Rejected('자신의 권한·활성 상태는 이 화면에서 변경할 수 없습니다.')
                        user.is_active=request.POST.get('active')=='on'
                        user.is_superuser=request.POST.get('role')=='admin';user.save()
                        svc.audit(request.user,'계정 권한 변경',user.username,active=user.is_active,admin=user.is_superuser)
                    else:
                        raise svc.Rejected('올바르지 않은 요청입니다.')
            messages.success(request,'계정을 저장했습니다.')
            return redirect('accounts')
        except (svc.Rejected,ValidationError) as e:
            messages.error(request,' '.join(e.messages) if isinstance(e,ValidationError) else str(e))
    return render(request,'accounts.html',{'accounts':User.objects.filter(is_staff=True).order_by('username')})

@staff_required
def sandbox(request):
    if not settings.SANDBOX:
        raise Http404()
    p=Policy.current()
    if request.method=='POST':
        action=request.POST.get('action','')
        try:
            if action=='reset':
                if request.POST.get('confirm')!='RESET':
                    raise svc.Rejected('초기화하려면 RESET을 입력해주세요.')
                svc.reset_sandbox(request.user)
            elif action=='seed':
                if Course.objects.exists():
                    raise svc.Rejected('기존 데이터가 있습니다. 전체 초기화를 사용해주세요.')
                svc.reset_sandbox(request.user)
            elif action=='location':
                mode=request.POST.get('mode')
                if mode not in dict(Policy._meta.get_field('location_mode').choices):
                    raise svc.Rejected('위치 모드를 선택해주세요.')
                p.location_mode=mode;p.save();svc.audit(request.user,'테스트 위치 변경',mode)
            else:
                c=Course.objects.first()
                if not c:
                    raise svc.Rejected('먼저 기본 테스트를 준비해주세요.')
                offsets={'normal':0,'before':-11*60,'late':16*60,'ended':61*60,'boundary':14*60+50,'closing':59*60+50}
                with transaction.atomic():
                    p=Policy.objects.select_for_update().get(pk=1)
                    if action in offsets:
                        p.generation+=1
                        svc.set_test_time(p,c.starts_at+timedelta(seconds=offsets[action]))
                    elif action=='forward':
                        svc.set_test_time(p,svc.now(p)+timedelta(minutes=1))
                    elif action=='expire':
                        Attempt.objects.filter(completed__isnull=True).update(expires_at=svc.now(p)-timedelta(seconds=1))
                    elif action=='duplicate':
                        e=c.enrollments.first()
                        Attendance.objects.get_or_create(enrollment=e,defaults={'status':'present','attended_at':c.starts_at,'method':'manual','reason':'중복 제출 시나리오','location_source':'mock-seed'})
                    else:
                        raise svc.Rejected('지원하지 않는 시나리오입니다.')
                    svc.audit(request.user,'테스트 시나리오',action)
            messages.success(request,'테스트 상황을 준비했습니다. 시간 시나리오를 바꿨다면 새 QR로 접속해주세요.')
            return redirect('sandbox')
        except svc.Rejected as e:
            messages.error(request,str(e))
    p=Policy.current()
    return render(request,'sandbox.html',{'policy':p,'clock':svc.now(p),'students':Enrollment.objects.filter(course=Course.objects.first())[:70],
        'scenarios':[('normal','정상 출석'),('before','접수 전'),('late','지각'),('ended','수업 종료'),('boundary','지각 전환 10초 전'),('closing','수업 종료 10초 전'),('forward','현재 시각 +1분'),('expire','작성 세션 즉시 만료'),('duplicate','첫 학생 중복 제출 준비')],
        'modes':Policy._meta.get_field('location_mode').choices})

def health(request):
    try:
        with connection.cursor() as cur:
            cur.execute('SELECT 1')
        return JsonResponse({'status':'ok'})
    except Exception:
        return JsonResponse({'status':'unavailable'},status=503)
