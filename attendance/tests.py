import base64
import io
import json
import tempfile
from datetime import timedelta
from unittest.mock import patch
from PIL import Image, ImageDraw
from django.conf import settings
from django.contrib.auth import get_user_model
from django.test import TestCase, Client, override_settings
from django.urls import reverse
from django.utils import timezone
from django.core.files.uploadedfile import SimpleUploadedFile
from .models import Policy, Course, Enrollment, Attendance, Attempt, Audit, BrowserIdentity
from . import services as svc

def signature():
    im=Image.new('RGB',(300,100),'white');draw=ImageDraw.Draw(im);draw.line([(15,70),(80,20),(130,70),(220,30)],fill='black',width=3)
    b=io.BytesIO();im.save(b,format='PNG')
    return 'data:image/png;base64,'+base64.b64encode(b.getvalue()).decode()

@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class AttendanceTests(TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.settings_override=override_settings(MEDIA_ROOT=self.tmp.name)
        self.settings_override.enable();self.addCleanup(self.settings_override.disable)
        self.p=Policy.current();self.p.latitude=35.1;self.p.longitude=129.0;self.p.save()
        self.at=timezone.now().replace(microsecond=0)
        self.c=Course.objects.create(title='수업',section='A',starts_at=self.at,ends_at=self.at+timedelta(hours=1))
        self.e=Enrollment.objects.create(course=self.c,department='학과',student_id='001234',name='홍길동')
        self.admin=get_user_model().objects.create_user('admin',password='Long-password-2026!',is_staff=True,is_superuser=True)
        self.staff=get_user_model().objects.create_user('assistant',password='Long-password-2026!',is_staff=True)

    def payload(self):
        return {'course':self.c.pk,'student_id':self.e.student_id,'name':self.e.name,'signature':signature(),
                'latitude':35.1,'longitude':129.0,'accuracy':20,'location_timestamp':timezone.now().timestamp()*1000}
    def attempt(self):
        return svc.exchange_qr()
    def attach_attempt(self,attempt):
        session=self.client.session;session['attempt']=str(attempt.pk);session.save()
        from django.http import HttpResponse
        response=HttpResponse()
        response.set_signed_cookie(settings.BROWSER_COOKIE_NAME,str(attempt.browser_id),salt='browser-identity')
        self.client.cookies[settings.BROWSER_COOKIE_NAME]=response.cookies[settings.BROWSER_COOKIE_NAME]

    def enter_browser(self,client):
        response=client.get(reverse('enter'))
        self.assertEqual(response.status_code,302)
        return Attempt.objects.get(pk=client.session['attempt'])

    def browser_submit(self,client,payload=None):
        return client.post(reverse('submit'),json.dumps(payload or self.payload()),content_type='application/json')

    def test_browser_binding_blocks_other_student_but_allows_other_course(self):
        a=self.enter_browser(self.client)
        self.assertTrue(self.client.cookies[settings.BROWSER_COOKIE_NAME]['httponly'])
        self.assertEqual(self.browser_submit(self.client).status_code,200)
        a.browser.refresh_from_db();self.assertEqual(a.browser.student_id,self.e.student_id)
        Enrollment.objects.create(course=self.c,student_id='999',name='다른학생')
        second=Course.objects.create(title='다음 수업',starts_at=self.c.starts_at,ends_at=self.c.ends_at)
        Enrollment.objects.create(course=second,student_id=self.e.student_id,name=self.e.name)
        b=self.enter_browser(self.client);self.assertEqual(b.browser_id,a.browser_id)
        response=self.client.get(reverse('checkin'));self.assertContains(response,'readonly')
        payload=self.payload();payload.update(student_id='999',name='다른학생')
        rejected=self.browser_submit(self.client,payload)
        self.assertEqual(rejected.status_code,400)
        self.assertIn('브라우저',rejected.json()['error'])
        self.assertEqual(Attendance.objects.count(),1)
        payload=self.payload();payload['course']=second.pk
        self.assertEqual(self.browser_submit(self.client,payload).status_code,200)
        self.assertEqual(Attendance.objects.count(),2)

    def test_failed_submit_does_not_bind_and_cookie_required(self):
        a=self.enter_browser(self.client)
        payload=self.payload();payload['signature']=''
        self.assertEqual(self.browser_submit(self.client,payload).status_code,400)
        a.browser.refresh_from_db();self.assertEqual(a.browser.student_id,'')
        cookie=self.client.cookies[settings.BROWSER_COOKIE_NAME].value
        self.client.cookies[settings.BROWSER_COOKIE_NAME]=cookie+'tampered'
        self.assertEqual(self.browser_submit(self.client).status_code,400)
        del self.client.cookies[settings.BROWSER_COOKIE_NAME]
        self.assertEqual(self.browser_submit(self.client).status_code,400)
        self.assertFalse(Attendance.objects.exists())

    def test_preopened_attempt_uses_latest_binding_and_other_browser_is_independent(self):
        first=self.enter_browser(self.client)
        preopened=self.enter_browser(self.client)
        svc.submit(first.pk,self.payload())
        Enrollment.objects.create(course=self.c,student_id='999',name='다른학생')
        payload=self.payload();payload.update(student_id='999',name='다른학생')
        with self.assertRaisesMessage(svc.Rejected,'브라우저'):
            svc.submit(preopened.pk,payload)
        other=Client();other_attempt=self.enter_browser(other)
        self.assertNotEqual(other_attempt.browser_id,first.browser_id)
        self.assertEqual(self.browser_submit(other,payload).status_code,200)

    def test_binding_survives_staff_logout_and_new_session(self):
        first=self.enter_browser(self.client)
        self.assertEqual(self.browser_submit(self.client).status_code,200)
        self.client.force_login(self.staff)
        self.client.post(reverse('logout'))
        second=self.enter_browser(self.client)
        self.assertEqual(first.browser_id,second.browser_id)
        self.assertEqual(second.browser.student_id,self.e.student_id)

    def test_new_browser_session_with_persistent_cookie_keeps_student_binding(self):
        first=self.enter_browser(self.client)
        self.assertEqual(self.browser_submit(self.client).status_code,200)
        restarted=Client()
        restarted.cookies[settings.BROWSER_COOKIE_NAME]=self.client.cookies[settings.BROWSER_COOKIE_NAME].value
        second=self.enter_browser(restarted)
        self.assertEqual(first.browser_id,second.browser_id)
        self.assertEqual(second.browser.student_id,self.e.student_id)
        meta=Audit.objects.get(action='QR 브라우저 확인',target=str(second.pk)).details
        self.assertEqual(meta['cookie_state'],'valid');self.assertTrue(meta['reused'])
        self.assertFalse(meta['session_cookie_present'])
        payload=self.payload();payload.update(student_id='999',name='다른학생')
        self.assertEqual(self.browser_submit(restarted,payload).status_code,400)

    def test_diagnostics_distinguish_missing_and_invalid_cookies_without_secrets(self):
        first=self.enter_browser(self.client)
        self.browser_submit(self.client)
        del self.client.cookies[settings.BROWSER_COOKIE_NAME]
        second=self.enter_browser(self.client)
        meta=Audit.objects.get(action='QR 브라우저 확인',target=str(second.pk)).details
        self.assertEqual(meta['cookie_state'],'missing')
        self.assertFalse(meta['reused']);self.assertTrue(meta['session_cookie_present'])
        self.assertEqual(meta['previous_browser_code'],first.browser.support_code)
        self.assertNotEqual(second.browser_id,first.browser_id)
        raw=self.client.cookies[settings.BROWSER_COOKIE_NAME].value+'tampered'
        self.client.cookies[settings.BROWSER_COOKIE_NAME]=raw
        third=self.enter_browser(self.client)
        meta=Audit.objects.get(action='QR 브라우저 확인',target=str(third.pk)).details
        self.assertEqual(meta['cookie_state'],'invalid_signature')
        self.assertNotIn(raw,json.dumps(meta))
        self.assertNotIn('ip',meta);self.assertNotIn('token',meta)

    def test_diagnostic_page_staff_only_and_old_records_remain_visible(self):
        a=self.attempt();svc.submit(a.pk,self.payload())
        url=reverse('browser_diagnostics')
        self.assertEqual(self.client.get(url).status_code,302)
        user=get_user_model().objects.create_user('ordinary',password='Long-password-2026!')
        self.client.force_login(user)
        self.assertEqual(self.client.get(url).status_code,403)
        self.client.force_login(self.staff)
        response=self.client.get(url)
        self.assertContains(response,a.browser.support_code)
        self.assertContains(response,'진단 추가 전 기록')
        self.assertContains(response,self.e.student_id)
        self.assertEqual(Attendance.objects.count(),1)
        self.assertEqual(self.client.post(url,{}).status_code,405)

    def test_expired_cookie_diagnostic(self):
        self.enter_browser(self.client)
        with patch('django.core.signing.time.time',return_value=timezone.now().timestamp()+settings.BROWSER_COOKIE_AGE+10):
            attempt=self.enter_browser(self.client)
        self.assertEqual(Audit.objects.get(action='QR 브라우저 확인',target=str(attempt.pk)).details['cookie_state'],'expired')

    def test_sandbox_reset_removes_bindings_but_scenario_changes_keep_them(self):
        with override_settings(SANDBOX=True):
            a=self.enter_browser(self.client);self.browser_submit(self.client)
            staff=Client();staff.force_login(self.staff)
            staff.post(reverse('sandbox'),{'action':'normal'})
            self.assertEqual(BrowserIdentity.objects.get(pk=a.browser_id).student_id,self.e.student_id)
            self.assertEqual(self.enter_browser(self.client).browser_id,a.browser_id)
            svc.reset_sandbox(self.admin)
            self.assertFalse(BrowserIdentity.objects.exists())
            fresh=self.enter_browser(self.client)
            self.assertNotEqual(fresh.browser_id,a.browser_id)
            self.assertEqual(fresh.browser.student_id,'')

    def test_student_has_no_identity_change_and_staff_handles_exception_manually(self):
        a=self.enter_browser(self.client)
        self.assertNotContains(self.client.get(reverse('checkin')),'학번 변경')
        self.browser_submit(self.client)
        self.assertNotContains(self.client.get(reverse('complete')),'학번 변경')
        row=Enrollment.objects.create(course=self.c,student_id='999',name='다른학생')
        staff=Client();staff.force_login(self.staff)
        self.assertEqual(staff.get(settings.URL_PREFIX+'/staff/browser-identity/').status_code,404)
        response=staff.post(reverse('record',args=[row.pk]),{'status':'present',
            'attended_at':timezone.localtime(self.at).strftime('%Y-%m-%dT%H:%M'),
            'reason':'현장에서 본인 확인 후 수동 출석','revision':'new'})
        self.assertEqual(response.status_code,302)
        record=Attendance.objects.get(enrollment=row)
        self.assertEqual(record.method,'manual');self.assertFalse(record.signature)
        a.browser.refresh_from_db();self.assertEqual(a.browser.student_id,self.e.student_id)
        self.assertTrue(Audit.objects.filter(actor=self.staff,action='출석 정정',target=str(row.pk)).exists())

    def test_course_delete_removes_only_target_records_and_signatures_after_commit(self):
        a=self.attempt();record=svc.submit(a.pk,self.payload())
        path=record.signature.name;storage=record.signature.storage
        other=Course.objects.create(title='보존 수업',starts_at=self.c.starts_at,ends_at=self.c.ends_at)
        other_row=Enrollment.objects.create(course=other,student_id='999',name='다른학생')
        payload=self.payload();payload.update(course=other.pk,student_id='999',name='다른학생')
        other_record=svc.submit(self.attempt().pk,payload)
        self.client.force_login(self.admin)
        url=reverse('course_delete',args=[self.c.pk])
        preview=self.client.get(url)
        self.assertEqual(preview.context['preview']['signatures'],1)
        self.assertTrue(Course.objects.filter(pk=self.c.pk).exists())
        with self.captureOnCommitCallbacks(execute=True):
            response=self.client.post(url,{'token':preview.context['token'],'confirm':self.c.title,'reason':'잘못 등록한 수업 삭제'})
            self.assertEqual(response.status_code,302)
            self.assertTrue(storage.exists(path))
        self.assertFalse(storage.exists(path))
        self.assertFalse(Course.objects.filter(pk=self.c.pk).exists())
        self.assertFalse(Enrollment.objects.filter(pk=self.e.pk).exists())
        self.assertFalse(Attendance.objects.filter(pk=record.pk).exists())
        self.assertFalse(Attempt.objects.filter(pk=a.pk).exists())
        self.assertTrue(Attendance.objects.filter(pk=other_record.pk).exists())
        self.assertTrue(other_record.signature.storage.exists(other_record.signature.name))
        self.assertTrue(BrowserIdentity.objects.filter(pk=a.browser_id).exists())
        log=Audit.objects.get(action='수업 삭제',target=str(self.c.pk))
        self.assertEqual(log.details['signatures'],1);self.assertEqual(log.actor,self.admin)
        self.assertEqual(self.client.post(url,{}).status_code,404)

    def test_course_delete_requires_staff_csrf_confirmation_reason_and_fresh_preview(self):
        url=reverse('course_delete',args=[self.c.pk])
        self.assertEqual(self.client.get(url).status_code,302)
        self.assertEqual(self.client.post(url,{}).status_code,302)
        self.client.force_login(self.staff)
        token=self.client.get(url).context['token']
        payload={'token':token,'confirm':self.c.title,'reason':'등록 실수'}
        strict=Client(enforce_csrf_checks=True);strict.force_login(self.admin)
        self.assertEqual(strict.post(url,payload).status_code,403)
        for update in [{'token':'bad'},{'confirm':'다른 수업'},{'reason':''}]:
            self.assertEqual(self.client.post(url,{**payload,**update}).status_code,200)
            self.assertTrue(Course.objects.filter(pk=self.c.pk).exists())
        svc.submit(self.attempt().pk,self.payload())
        self.assertContains(self.client.post(url,payload),'삭제 범위를 다시 확인')
        self.assertTrue(Course.objects.filter(pk=self.c.pk).exists())

    def test_assistant_can_delete_empty_course_and_rollback_preserves_files(self):
        from django.core import signing
        record=svc.submit(self.attempt().pk,self.payload())
        token=signing.dumps(svc.course_delete_preview(self.c),salt='course-delete')
        with patch('attendance.services.audit',side_effect=RuntimeError('audit unavailable')):
            with self.assertRaises(RuntimeError):
                svc.delete_course(self.c.pk,token,self.c.title,'실수',self.staff)
        self.assertTrue(Course.objects.filter(pk=self.c.pk).exists())
        self.assertTrue(record.signature.storage.exists(record.signature.name))
        empty=Course.objects.create(title='빈 수업',starts_at=self.c.starts_at,ends_at=self.c.ends_at)
        self.client.force_login(self.staff)
        url=reverse('course_delete',args=[empty.pk])
        token=self.client.get(url).context['token']
        self.assertEqual(self.client.post(url,{'token':token,'confirm':empty.title,'reason':'중복 등록'}).status_code,302)
        self.assertFalse(Course.objects.filter(pk=empty.pk).exists())

    def test_time_boundaries(self):
        cases=[(-601,None),(-600,'present'),(900,'present'),(901,'late'),(3599,'late'),(3600,None)]
        for seconds,expected in cases:
            with self.subTest(seconds=seconds):
                if expected is None:
                    with self.assertRaises(svc.Rejected):svc.course_status(self.c,self.p,self.at+timedelta(seconds=seconds))
                else:self.assertEqual(svc.course_status(self.c,self.p,self.at+timedelta(seconds=seconds)),expected)

    def test_valid_submit_and_idempotent_retry(self):
        a=self.attempt();record=svc.submit(a.pk,self.payload())
        same=svc.submit(a.pk,self.payload())
        self.assertEqual(record.pk,same.pk);self.assertEqual(Attendance.objects.count(),1)
        self.assertTrue(record.signature.storage.exists(record.signature.name))
        self.assertEqual(record.status,'present')

    def test_different_session_cannot_disclose_or_overwrite_existing_record(self):
        record=svc.submit(self.attempt().pk,self.payload())
        with self.assertRaisesMessage(svc.Rejected,'이미 처리된'):svc.submit(self.attempt().pk,self.payload())
        record.refresh_from_db();self.assertEqual(record.status,'present')

    def test_fixed_qr_and_download_stay_identical_after_time_and_reset(self):
        self.client.force_login(self.staff)
        original=self.client.get(reverse('qr_data')).json()
        self.assertEqual(original['url'],settings.PUBLIC_ORIGIN+reverse('enter'))
        with override_settings(SANDBOX=True):
            svc.reset_sandbox(self.admin)
            self.p.refresh_from_db()
            svc.set_test_time(self.p,svc.now(self.p)+timedelta(days=3))
            self.assertEqual(self.client.get(reverse('qr_data')).json(),original)
        response=self.client.get(reverse('qr_data'),{'download':'1'})
        self.assertEqual(response['Content-Type'],'image/png')
        self.assertTrue(b''.join(response.streaming_content).startswith(b'\x89PNG'))
        self.assertEqual(self.client.get(reverse('enter')).status_code,302)

    def test_writing_expired(self):
        a=self.attempt();a.expires_at=svc.now()-timedelta(seconds=1);a.save()
        with self.assertRaises(svc.Rejected):svc.submit(a.pk,self.payload())
        self.assertFalse(Attendance.objects.exists())

    def test_submission_after_class_end_even_with_valid_attempt(self):
        a=self.attempt();self.c.ends_at=self.at-timedelta(seconds=1);self.c.starts_at=self.at-timedelta(hours=1);self.c.save()
        with self.assertRaises(svc.Rejected):svc.submit(a.pk,self.payload())

    def test_wrong_name_and_course_rejected(self):
        payload=self.payload();payload['name']='다른학생'
        with self.assertRaises(svc.Rejected):svc.submit(self.attempt().pk,payload)

    def test_location_rejections(self):
        for change in [{'latitude':36.1},{'accuracy':900},{'latitude':float('nan')},{'longitude':190},{'accuracy':-1},{'location_timestamp':0}]:
            with self.subTest(change=change):
                payload=self.payload();payload.update(change)
                with self.assertRaises(svc.Rejected):svc.location_check(self.p,payload)

    def test_live_ignores_mock_location_and_test_clock(self):
        with override_settings(SANDBOX=False):
            self.p.location_mode='inside';self.p.test_time=self.at-timedelta(days=5);self.p.test_reference=timezone.now();self.p.save()
            self.assertLess(abs((svc.now(self.p)-timezone.now()).total_seconds()),2)
            with self.assertRaises(svc.Rejected):svc.location_check(self.p,{})
            with self.assertRaises(svc.Rejected):svc.reset_sandbox()

    def test_blank_or_invalid_signature(self):
        for data in ['', 'data:image/svg+xml;base64,AAAA','data:image/png;base64,???']:
            with self.assertRaises(svc.Rejected):svc.signature_image(data)
        im=Image.new('RGB',(300,100),'white');out=io.BytesIO();im.save(out,format='PNG')
        with self.assertRaises(svc.Rejected):svc.signature_image('data:image/png;base64,'+base64.b64encode(out.getvalue()).decode())

    def test_sandbox_reset_revokes_sessions_and_preserves_accounts(self):
        with override_settings(SANDBOX=True):
            record=svc.submit(self.attempt().pk,self.payload());path=record.signature.name;storage=record.signature.storage
            old=self.attempt()
            with self.captureOnCommitCallbacks(execute=True):svc.reset_sandbox(self.admin)
            self.assertEqual(Course.objects.count(),3);self.assertEqual(Enrollment.objects.count(),210)
            self.assertFalse(Attendance.objects.exists());self.assertFalse(Attempt.objects.exists())
            self.assertFalse(storage.exists(path));self.assertTrue(get_user_model().objects.filter(pk=self.admin.pk).exists())
            self.assertFalse(Attempt.objects.filter(pk=old.pk).exists())
            self.assertEqual(svc.exchange_qr().generation,Policy.current().generation)

    def test_mock_location_is_explicit_in_record(self):
        with override_settings(SANDBOX=True):
            self.p.location_mode='inside';self.p.save()
            payload=self.payload()
            record=svc.submit(self.attempt().pk,payload)
            self.assertEqual(record.location_source,'mock-inside');self.assertEqual(record.latitude,35.1)

    def test_student_cannot_read_private_pages_or_signature(self):
        record=svc.submit(self.attempt().pk,self.payload())
        for name,args in [('dashboard',[]),('accounts',[]),('export',[]),('signature',[record.pk]),('course',[self.c.pk])]:
            self.assertEqual(self.client.get(reverse(name,args=args)).status_code,302)

    def test_staff_cannot_administer_system(self):
        self.client.force_login(self.staff)
        for name in ['accounts','system']:
            self.assertEqual(self.client.get(reverse(name)).status_code,403)
            self.assertEqual(self.client.post(reverse(name),{}).status_code,403)

    def test_all_staff_screens_render(self):
        self.client.force_login(self.admin)
        for name,args in [('dashboard',[]),('courses',[]),('course_new',[]),('course_edit',[self.c.pk]),('course',[self.c.pk]),('record',[self.e.pk]),('import',[self.c.pk]),('location',[]),('qr',[]),('qr_data',[]),('accounts',[]),('system',[]),('audit',[])]:
            with self.subTest(name=name):self.assertEqual(self.client.get(reverse(name,args=args)).status_code,200)
        if settings.SANDBOX:self.assertEqual(self.client.get(reverse('sandbox')).status_code,200)
        else:self.assertEqual(self.client.get(reverse('sandbox')).status_code,404)

    def test_checkin_and_completed_flow(self):
        self.attach_attempt(self.attempt())
        response=self.client.get(reverse('checkin'));self.assertEqual(response.status_code,200)
        response=self.client.post(reverse('submit'),data=json.dumps(self.payload()),content_type='application/json')
        self.assertEqual(response.status_code,200)
        response=self.client.get(reverse('complete'));self.assertContains(response,'출석 완료')

    def test_submission_requires_csrf(self):
        client=Client(enforce_csrf_checks=True)
        response=client.post(reverse('submit'),data=json.dumps(self.payload()),content_type='application/json')
        self.assertEqual(response.status_code,403)

    def test_login_csrf_origin_and_referrer_policy(self):
        client=Client(enforce_csrf_checks=True)
        response=client.get(reverse('login'))
        self.assertEqual(response['Referrer-Policy'],'same-origin')
        csrf=client.cookies[settings.CSRF_COOKIE_NAME].value
        response=client.post(reverse('login'),{'username':'assistant','password':'Long-password-2026!','csrfmiddlewaretoken':csrf},HTTP_ORIGIN='http://testserver')
        self.assertEqual(response.status_code,302)
        self.assertEqual(client.get(reverse('dashboard')).status_code,200)

    def test_export_keeps_leading_zero_and_embeds_signature(self):
        svc.submit(self.attempt().pk,self.payload());self.client.force_login(self.staff)
        import openpyxl
        for with_signature in [False,True]:
            response=self.client.get(reverse('export'),{'signatures':'1' if with_signature else '0'})
            workbook=openpyxl.load_workbook(io.BytesIO(b''.join(response.streaming_content)))
            sheet=workbook.active
            self.assertEqual(sheet['E2'].value,'001234');self.assertEqual(len(sheet._images),1 if with_signature else 0)

    def test_roster_import_preview_and_preserve_existing(self):
        import openpyxl
        book=openpyxl.Workbook();book.active.append(['학과','학번','이름']);book.active.append(['테스트','000001','학생1'])
        out=io.BytesIO();book.save(out);self.client.force_login(self.staff)
        url=reverse('import',args=[self.c.pk])
        r=self.client.post(url,{'action':'upload','file':SimpleUploadedFile('roster.xlsx',out.getvalue())});self.assertEqual(r.status_code,200)
        r=self.client.post(url,{'action':'map','department':'0','student_id':'1','name':'2'});self.assertContains(r,'000001')
        self.assertFalse(Enrollment.objects.filter(student_id='000001').exists())
        r=self.client.post(url,{'action':'confirm'});self.assertEqual(r.status_code,302)
        self.assertTrue(Enrollment.objects.filter(pk=self.e.pk).exists());self.assertTrue(Enrollment.objects.filter(student_id='000001').exists())

    def test_formula_roster_rejected(self):
        import openpyxl
        from .views import read_xlsx
        book=openpyxl.Workbook();book.active.append(['학과','학번','이름']);book.active.append(['a','=1+1','b'])
        out=io.BytesIO();book.save(out)
        with self.assertRaises(svc.Rejected):read_xlsx(SimpleUploadedFile('bad.xlsx',out.getvalue()))

    def test_manual_correction_audit_and_stale_write(self):
        from .views import revision
        self.client.force_login(self.staff)
        url=reverse('record',args=[self.e.pk])
        payload={'status':'late','attended_at':timezone.localtime(self.at).strftime('%Y-%m-%dT%H:%M'),'reason':'위치 오류 현장 확인','revision':'new'}
        self.assertEqual(self.client.post(url,payload).status_code,302)
        record=Attendance.objects.get(enrollment=self.e);self.assertEqual(record.method,'manual');self.assertFalse(record.signature)
        payload['status']='absent'
        self.assertContains(self.client.post(url,payload),'다른 작업에서 기록이 변경')
        record.refresh_from_db();self.assertEqual(record.status,'late');self.assertTrue(Audit.objects.filter(action='출석 정정').exists())

    def test_test_buttons_and_expiration(self):
        with override_settings(SANDBOX=True):
            self.client.force_login(self.staff)
            r=self.client.post(reverse('sandbox'),{'action':'late'});self.assertEqual(r.status_code,302)
            self.p.refresh_from_db();self.assertEqual(svc.course_status(self.c,self.p),'late')
            attempt=self.attempt()
            self.client.post(reverse('sandbox'),{'action':'expire'})
            attempt.refresh_from_db()
            with self.assertRaises(svc.Rejected):svc.valid_attempt(attempt,self.p)

    def test_unsafe_spreadsheet_cells_are_neutralized(self):
        from .views import safe_cell
        self.assertEqual(safe_cell('=HYPERLINK("bad")'),'\'=HYPERLINK("bad")')
        self.assertEqual(safe_cell('0012'),'0012')

    def test_anonymous_status_never_returns_roster(self):
        response=self.client.get(reverse('attempt_status'))
        self.assertNotIn('name',response.json());self.assertNotIn('student_id',response.json())

    def guest_payload(self):
        return {**self.payload(),'kind':'guest','department':'다른학과','student_id':'000987','name':'방문학생'}

    def test_guest_daily_limit_and_next_day_without_roster(self):
        self.enter_browser(self.client)
        payload=self.guest_payload()
        self.assertEqual(self.browser_submit(self.client,payload).status_code,200)
        record=Attendance.objects.get(enrollment__isnull=True)
        self.assertEqual(record.guest_date,timezone.localtime(svc.now()).date())
        self.assertEqual(record.student_id,'000987');self.assertEqual(record.status,'present')
        self.assertEqual(Enrollment.objects.count(),1)
        self.assertContains(self.client.get(reverse('complete')),'다른학과')
        self.assertNotContains(self.client.get(reverse('complete')),'<dt>수업</dt>')
        self.assertEqual(self.browser_submit(self.client,payload).status_code,200) # same-attempt retry
        self.enter_browser(self.client)
        self.assertIn('하루에 한 번',self.browser_submit(self.client,payload).json()['error'])
        with patch('attendance.services.now',return_value=self.at+timedelta(days=1)):
            self.enter_browser(self.client)
            self.assertEqual(self.browser_submit(self.client,payload).status_code,200)
        self.assertEqual(Attendance.objects.filter(enrollment__isnull=True).count(),2)

    def test_guest_validations_and_browser_binding_shared_with_courses(self):
        self.enter_browser(self.client)
        payload=self.guest_payload()
        for invalid in [{'department':''},{'name':''},{'student_id':'x'*41},{'signature':''},{'latitude':None},{'kind':'unknown'}]:
            self.assertEqual(self.browser_submit(self.client,{**payload,**invalid}).status_code,400)
        payload.update(student_id=self.e.student_id,name=self.e.name)
        self.assertEqual(self.browser_submit(self.client,payload).status_code,200)
        self.enter_browser(self.client)
        self.assertEqual(self.browser_submit(self.client,{**payload,'student_id':'other'}).status_code,400)
        self.assertEqual(self.browser_submit(self.client).status_code,200)
        self.assertEqual(Attendance.objects.count(),2)

    def test_location_required_even_in_all_sandbox_modes(self):
        with override_settings(SANDBOX=True):
            for mode in ['real','inside','outside','inaccurate','unavailable']:
                self.p.location_mode=mode;self.p.save()
                for kind in ['course','guest']:
                    payload=self.guest_payload() if kind=='guest' else self.payload()
                    payload.pop('latitude')
                    with self.assertRaisesMessage(svc.Rejected,'위치 정보가 없습니다'):
                        svc.submit(self.attempt().pk,payload)
            self.assertFalse(Attendance.objects.exists())

    def test_only_current_courses_visible_and_status_updates_at_boundary(self):
        self.enter_browser(self.client)
        Course.objects.create(title='미래수업',starts_at=self.at+timedelta(minutes=11),ends_at=self.at+timedelta(hours=2))
        Course.objects.create(title='종료수업',starts_at=self.at-timedelta(hours=1),ends_at=self.at)
        Course.objects.create(title='중지수업',starts_at=self.at,ends_at=self.at+timedelta(hours=1),active=False)
        with patch('attendance.services.now',return_value=self.at):
            response=self.client.get(reverse('checkin'))
            for title in ['미래수업','종료수업','중지수업']:self.assertNotContains(response,title)
            self.assertEqual([c['id'] for c in self.client.get(reverse('attempt_status')).json()['courses']],[self.c.pk])
        with patch('attendance.services.now',return_value=self.at+timedelta(hours=2)):
            data=self.client.get(reverse('attempt_status')).json()
            self.assertEqual(data['courses'],[])
            self.assertNotIn(str(self.e.student_id),json.dumps(data))

    def test_server_clock_distinguishes_actual_and_test_time(self):
        with override_settings(SANDBOX=True):
            svc.set_test_time(self.p,self.at-timedelta(days=3))
            data=self.client.get(reverse('clock')).json()
            from datetime import datetime
            self.assertLess(abs((datetime.fromisoformat(data['server_time'])-timezone.now()).total_seconds()),3)
            self.assertGreater((datetime.fromisoformat(data['server_time'])-datetime.fromisoformat(data['attendance_time'])).total_seconds(),2*86400)
            self.assertContains(self.client.get(reverse('home')),'서버 시각')

    def test_guest_staff_permissions_exports_and_manual_edit(self):
        record=svc.submit(self.attempt().pk,self.guest_payload())
        for route in ['guests','guest_export','guest_new']:
            self.assertEqual(self.client.get(reverse(route)).status_code,302)
        self.assertEqual(self.client.get(reverse('guest_record',args=[record.pk])).status_code,302)
        self.client.force_login(self.staff)
        self.assertContains(self.client.get(reverse('guests')),'방문학생')
        self.assertContains(self.client.get(reverse('browser_diagnostics')),'000987')
        import openpyxl
        for signatures in ['0','1']:
            r=self.client.get(reverse('guest_export'),{'signatures':signatures})
            sheet=openpyxl.load_workbook(io.BytesIO(b''.join(r.streaming_content))).active
            self.assertEqual(sheet['C2'].value,'000987');self.assertEqual(len(sheet._images),int(signatures))
            self.assertNotIn('수업명',[c.value for c in sheet[1]])
        r=self.client.get(reverse('export'))
        sheet=openpyxl.load_workbook(io.BytesIO(b''.join(r.streaming_content))).active
        self.assertEqual(sheet.max_row,2);self.assertEqual(sheet['E2'].value,self.e.student_id)
        url=reverse('guest_record',args=[record.pk]);page=self.client.get(url)
        data={**page.context['form'].initial,'name':'수정이름','reason':'현장에서 본인 확인'}
        self.assertEqual(self.client.post(url,data).status_code,302)
        record.refresh_from_db();self.assertEqual(record.name,'수정이름');self.assertEqual(record.signed_name,'방문학생')
        self.assertTrue(record.signature);self.assertEqual(record.method,'manual')
        self.assertContains(self.client.post(url,data),'기록이 변경')
        self.assertTrue(Audit.objects.filter(action='타 학과 출석 정정').exists())

    def test_guest_reset_removes_signature_and_courses_delete_preserves_guest(self):
        guest=svc.submit(self.attempt().pk,self.guest_payload());path=guest.signature.name;storage=guest.signature.storage
        self.client.force_login(self.staff)
        url=reverse('course_delete',args=[self.c.pk]);token=self.client.get(url).context['token']
        self.client.post(url,{'token':token,'confirm':self.c.title,'reason':'잘못 등록한 수업'})
        self.assertTrue(Attendance.objects.filter(pk=guest.pk).exists())
        with override_settings(SANDBOX=True),self.captureOnCommitCallbacks(execute=True):
            svc.reset_sandbox(self.staff)
        self.assertFalse(Attendance.objects.exists());self.assertFalse(storage.exists(path))

    def test_guest_duplicate_other_browser_and_korean_midnight(self):
        from datetime import datetime, timezone as dt_timezone
        payload=self.guest_payload()
        with patch('attendance.services.now',return_value=datetime(2026,10,25,14,59,59,tzinfo=dt_timezone.utc)):
            first=svc.submit(self.attempt().pk,payload)
            self.assertEqual(first.guest_date.isoformat(),'2026-10-25')
            with self.assertRaisesMessage(svc.Rejected,'하루에 한 번'):
                svc.submit(self.attempt().pk,payload)
        with patch('attendance.services.now',return_value=datetime(2026,10,25,15,0,0,tzinfo=dt_timezone.utc)):
            second=svc.submit(self.attempt().pk,payload)
            self.assertEqual(second.guest_date.isoformat(),'2026-10-26')

    def test_guest_manual_duplicate_and_database_constraint(self):
        record=svc.submit(self.attempt().pk,self.guest_payload())
        self.client.force_login(self.staff)
        data={'department':record.department,'student_id':record.student_id,'name':record.name,
            'guest_date':record.guest_date,'attended_at':record.attended_at,'status':'present','reason':'현장 확인 처리','revision':'new'}
        self.assertContains(self.client.post(reverse('guest_new'),data),'출석이 이미 있습니다')
        data['student_id']='new-student'
        self.assertEqual(self.client.post(reverse('guest_new'),data).status_code,302)
        self.assertEqual(Attendance.objects.filter(enrollment__isnull=True).count(),2)
        from django.db import IntegrityError,transaction
        with self.assertRaises(IntegrityError),transaction.atomic():
            Attendance.objects.create(guest_date=record.guest_date,student_id=record.student_id,department=record.department,name=record.name,status='present')


from django.test import TransactionTestCase
from django.db import connection
from django.db.migrations.executor import MigrationExecutor

class MigrationPreservationTests(TransactionTestCase):
    def test_existing_roster_and_signed_attendance_survive_upgrade(self):
        previous=[('attendance','0002_browseridentity_attempt_browser_revision_and_more')]
        latest=[('attendance','0003_remove_policy_qr_seconds_attendance_department_and_more')]
        executor=MigrationExecutor(connection)
        executor.migrate(previous)
        try:
            apps=executor.loader.project_state(previous).apps
            at=timezone.now()
            course=apps.get_model('attendance','Course').objects.create(title='기존 수업',starts_at=at,ends_at=at+timedelta(hours=1))
            row=apps.get_model('attendance','Enrollment').objects.create(course=course,student_id='001',name='기존 학생')
            saved=apps.get_model('attendance','Attendance').objects.create(enrollment=row,status='present',attended_at=at,signature='signatures/existing.png',signed_name='기존 학생')
            MigrationExecutor(connection).migrate(latest)
            record=Attendance.objects.get(pk=saved.pk)
            self.assertEqual(record.enrollment.student_id,'001')
            self.assertEqual(record.signature.name,'signatures/existing.png')
            self.assertEqual(record.signed_name,'기존 학생');self.assertIsNone(record.guest_date)
        finally:
            MigrationExecutor(connection).migrate(latest)
